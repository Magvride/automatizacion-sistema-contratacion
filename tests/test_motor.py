# -*- coding: utf-8 -*-
"""Test de integración del motor de auditoría (``src/alfresco_mcp/motor.py``)."""

import openpyxl

from alfresco_mcp.gateway import AlfrescoGateway
from alfresco_mcp.motor import ejecutar_auditoria


class FakeGateway(AlfrescoGateway):
    def __init__(self, por_numero):
        self._por_numero = por_numero

    def buscar_carpetas(self, numero, max_items=10):
        return self._por_numero.get(numero, {}).get("carpetas", [])

    def listar_archivos(self, node_id, max_items=100):
        for datos in self._por_numero.values():
            for carpeta in datos.get("carpetas", []):
                if carpeta["id"] == node_id:
                    archivos = datos.get("archivos", [])
                    return {"archivos": archivos, "total": len(archivos)}
        return {"archivos": [], "total": 0}

    def obtener_nodo(self, node_id):
        return None

    def descargar(self, node_id, destino):
        return destino


def _diccionario(tmp_path):
    ruta = tmp_path / "Diccionario_Documentos.xlsx"
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Diccionario"
    hoja.append(["ID_DOC", "CLASE", "CODIGO_CLASE", "DOCUMENTO", "CODIGO_FORMATO",
                 "ETAPA", "CARDINALIDAD", "OBLIGATORIEDAD", "ALIAS"])
    hoja.append(["D190", "ORDEN DE COMPRA", "20", "Informe de oportunidad",
                 "FCO.55", "Precontractual - selección", "1", "Obligatorio",
                 "FCO.55 Informe Oportunidad"])
    libro.save(ruta)
    return str(ruta)


def test_ejecutar_auditoria_end_to_end(tmp_path):
    gateway = FakeGateway(
        {
            "2026000003": {
                "carpetas": [{"id": "n1", "nombre": "0020_2026000003_9707", "es_carpeta": True}],
                "archivos": [{"NOMBRE": "FCO.55_Informe.pdf", "ID": "d1",
                              "FECHA_CREACION": "2026-01-01", "FECHA_MODIFICACION": ""}],
            }
        }
    )
    contratos = [
        {"contrato": "20-2026000003", "centro_costo": "CC1"},
        {"contrato": "20-2026009999", "centro_costo": "CC2"},
    ]
    progreso = []
    resumen = ejecutar_auditoria(
        contratos,
        gateway,
        ruta_diccionario=_diccionario(tmp_path),
        ruta_06=str(tmp_path / "06.csv"),
        ruta_07=str(tmp_path / "07.csv"),
        ruta_09=str(tmp_path / "09.xlsx"),
        periodo="PRUEBA",
        on_progreso=lambda hechos, total: progreso.append((hechos, total)),
    )

    assert resumen["total"] == 2
    assert resumen["encontradas"] == 1
    assert resumen["no_encontradas"] == 1
    assert progreso == [(1, 2), (2, 2)]

    libro = openpyxl.load_workbook(resumen["ruta_09"])
    assert "20_2026000003" in libro.sheetnames
    granular = libro["20_2026000003"]
    assert granular.cell(row=7, column=2).value == "D190"
    assert granular.cell(row=7, column=7).value == "ENCONTRADO"

    diagnostico = libro["DIAGNOSTICO"]
    from auditoria_documental.workbook import CABECERAS_DIAGNOSTICO

    def col(nombre):
        return CABECERAS_DIAGNOSTICO.index(nombre) + 1

    assert diagnostico.cell(row=5, column=col("CONTRATO")).value == "20-2026000003"
    assert diagnostico.cell(row=5, column=col("CLASE")).value == "20"
    assert diagnostico.cell(row=5, column=col("ESTADO")).value == "Completo"

    with open(resumen["ruta_07"], encoding="utf-8-sig") as archivo:
        contenido = archivo.read().strip().splitlines()
    assert len(contenido) == 2  # encabezado + 1 no encontrado
    assert "carpeta_no_encontrada" in contenido[1]


class ClonableFakeGateway(FakeGateway):
    """Fake con ``clonar`` para ejercitar la ruta en paralelo."""

    def clonar(self):
        return ClonableFakeGateway(self._por_numero)


def test_ejecutar_auditoria_paralelo(tmp_path):
    gateway = ClonableFakeGateway(
        {
            "2026000003": {
                "carpetas": [{"id": "n1", "nombre": "0020_2026000003_9707", "es_carpeta": True}],
                "archivos": [{"NOMBRE": "FCO.55_Informe.pdf", "ID": "d1",
                              "FECHA_CREACION": "", "FECHA_MODIFICACION": ""}],
            }
        }
    )
    contratos = [
        {"contrato": "20-2026000003"},
        {"contrato": "20-2026009001"},
        {"contrato": "20-2026009002"},
        {"contrato": "20-2026009003"},
    ]
    progreso = []
    resumen = ejecutar_auditoria(
        contratos,
        gateway,
        ruta_diccionario=_diccionario(tmp_path),
        ruta_06=str(tmp_path / "06.csv"),
        ruta_07=str(tmp_path / "07.csv"),
        ruta_09=str(tmp_path / "09.xlsx"),
        max_workers=3,
        on_progreso=lambda hechos, total: progreso.append((hechos, total)),
    )

    assert resumen["total"] == 4
    assert resumen["encontradas"] == 1
    assert resumen["no_encontradas"] == 3
    assert progreso[-1] == (4, 4)

    # El orden de los contratos se conserva pese al paralelismo.
    libro = openpyxl.load_workbook(resumen["ruta_09"])
    from auditoria_documental.workbook import CABECERAS_DIAGNOSTICO

    col = CABECERAS_DIAGNOSTICO.index("CONTRATO") + 1
    hoja = libro["DIAGNOSTICO"]
    orden = [hoja.cell(row=f, column=col).value for f in range(5, 9)]
    assert orden == [c["contrato"] for c in contratos]
