# -*- coding: utf-8 -*-
"""Tests del Excel de auditoría (``src/auditoria_documental/workbook.py``)."""

import openpyxl

from auditoria_documental.workbook import (
    CABECERAS_DIAGNOSTICO,
    CABECERAS_RESUMEN,
    generar_auditoria,
)


def _contrato(contrato, filas):
    return {
        "contrato": contrato,
        "cod": contrato.split("-")[0],
        "tipo": "Orden Compra",
        "contratista": "CASA HERMES LTDA",
        "centro_costo": "CC-1",
        "supervisor": "JUAN PEREZ",
        "carpeta": "0020_2026000003_9707",
        "node_id": "fecb2f27-d4d2-47c2-b66c-77c1e0b0b5a1",
        "estado_alfresco": "ENCONTRADA",
        "filas": filas,
        "etapas": {"Precontractual - selección": (1, 1)},
    }


def test_generar_auditoria_estructura(tmp_path):
    filas = [
        {"ID_DOC": "D190", "DOCUMENTO": "Informe oportunidad", "FORMATO": "FCO.55",
         "ETAPA": "Precontractual - selección", "ALIAS": "FCO.55", "ESTADO": "ENCONTRADO",
         "ARCHIVO": "FCO.55_Informe.pdf"},
        {"ID_DOC": "D027", "DOCUMENTO": "Solicitud cotizaciones", "FORMATO": "FCO.57",
         "ETAPA": "Precontractual - selección", "ALIAS": "FCO.57", "ESTADO": "FALTANTE",
         "ARCHIVO": ""},
    ]
    contratos = [
        _contrato("20-2026000003", filas),
        _contrato("18-2026001182", filas[:1]),
    ]
    salida = str(tmp_path / "09_Auditoria.xlsx")

    generar_auditoria(contratos, salida, periodo="ENERO 2026", notas=["Nota 1", "Nota 2"])

    libro = openpyxl.load_workbook(salida)
    assert libro.sheetnames == [
        "Resumen_Cruce_Corregido",
        "Cruce_por_Etapa_Corr",
        "DIAGNOSTICO",
        "20_2026000003",
        "18_2026001182",
        "Notas_Diccionario",
    ]

    resumen = libro["Resumen_Cruce_Corregido"]
    total_cols = len(CABECERAS_RESUMEN)
    assert [resumen.cell(row=4, column=c).value for c in range(1, total_cols + 1)] == CABECERAS_RESUMEN

    def col(nombre):
        return CABECERAS_RESUMEN.index(nombre) + 1

    assert resumen.cell(row=5, column=col("#")).value == 1
    assert resumen.cell(row=5, column=col("CÓDIGO CONTRATO")).value == "20-2026000003"
    assert resumen.cell(row=5, column=col("CENTRO DE COSTO")).value == "CC-1"
    assert resumen.cell(row=5, column=col("ORDENADOR")).value == "JUAN PEREZ"
    assert resumen.cell(row=5, column=col("ESPERADOS")).value == 2
    assert resumen.cell(row=5, column=col("ENCONTRADOS")).value == 1
    assert resumen.cell(row=5, column=col("FALTANTES")).value == 1
    assert resumen.cell(row=5, column=col("% CUMPL DIC")).value == 50.0
    assert resumen.cell(row=5, column=col("VÍNCULO CARPETA")).hyperlink is not None

    granular = libro["20_2026000003"]
    assert [granular.cell(row=5, column=c).value for c in range(1, 11)] == [
        "#", "ID_DOC", "DOCUMENTO", "FORMATO", "ETAPA", "ALIAS", "ESTADO",
        "ARCHIVO", "FECHA CREACIÓN/CARGA", "ÚLTIMA MODIFICACIÓN",
    ]
    assert granular.cell(row=6, column=1).value == "← Volver al Resumen"
    assert granular.cell(row=7, column=2).value == "D190"
    assert granular.cell(row=7, column=7).value == "ENCONTRADO"
    assert granular.cell(row=8, column=7).value == "FALTANTE"

    etapas = libro["Cruce_por_Etapa_Corr"]
    assert etapas.cell(row=4, column=1).value == "ETAPA"
    assert etapas.cell(row=5, column=2).value == "1/1"

    notas = libro["Notas_Diccionario"]
    assert notas.cell(row=3, column=1).value == "Nota 1"


def test_hoja_diagnostico(tmp_path):
    filas = [
        {"ID_DOC": "D190", "DOCUMENTO": "Informe oportunidad", "FORMATO": "FCO.55",
         "ETAPA": "Precontractual - selección", "ALIAS": "FCO.55", "ESTADO": "ENCONTRADO",
         "ARCHIVO": "FCO.55_Informe.pdf"},
        {"ID_DOC": "D027", "DOCUMENTO": "Solicitud cotizaciones", "FORMATO": "FCO.57",
         "ETAPA": "Precontractual - selección", "ALIAS": "FCO.57", "ESTADO": "FALTANTE",
         "ARCHIVO": ""},
    ]
    contratos = [_contrato("20-2026000003", filas)]
    salida = str(tmp_path / "09.xlsx")

    generar_auditoria(contratos, salida, periodo="ENERO 2026", fecha_revision="2026-08-31")

    libro = openpyxl.load_workbook(salida)
    ws = libro["DIAGNOSTICO"]
    total_cols = len(CABECERAS_DIAGNOSTICO)
    assert [ws.cell(row=4, column=c).value for c in range(1, total_cols + 1)] == CABECERAS_DIAGNOSTICO

    def col(nombre):
        return CABECERAS_DIAGNOSTICO.index(nombre) + 1

    assert ws.cell(row=5, column=col("CONTRATO")).value == "20-2026000003"
    assert ws.cell(row=5, column=col("CLASE")).value == "20"
    assert ws.cell(row=5, column=col("CARPETA ENCONTRADA")).value == "Encontrada"
    assert ws.cell(row=5, column=col("DOCS REQUERIDOS")).value == 2
    assert ws.cell(row=5, column=col("DOCS HALLADOS")).value == 1
    assert ws.cell(row=5, column=col("DOCS FALTANTES")).value == 1
    assert "FCO.57" in ws.cell(row=5, column=col("LISTADO DE FALTANTES")).value
    assert ws.cell(row=5, column=col("% COMPLETITUD")).value == "50.0%"
    assert ws.cell(row=5, column=col("ESTADO")).value == "Incompleto"
    assert ws.cell(row=5, column=col("FECHA REVISIÓN")).value == "2026-08-31"
    assert ws.cell(row=5, column=col("ENLACE CARPETA ALFRESCO")).hyperlink is not None


def test_nombre_hoja_unico_y_valido(tmp_path):
    filas = [{"ID_DOC": "D1", "DOCUMENTO": "x", "ESTADO": "ENCONTRADO"}]
    # Dos contratos con el mismo nombre generan hojas distintas.
    contratos = [_contrato("20-2026000003", filas), _contrato("20-2026000003", filas)]
    salida = str(tmp_path / "09.xlsx")
    generar_auditoria(contratos, salida)
    libro = openpyxl.load_workbook(salida)
    assert "20_2026000003" in libro.sheetnames
    assert len([n for n in libro.sheetnames if n.startswith("20_2026000003")]) == 2
