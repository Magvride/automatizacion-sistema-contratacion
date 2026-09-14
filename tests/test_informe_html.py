# -*- coding: utf-8 -*-
"""Tests del informe HTML (``src/auditoria_documental/informe_html.py``)."""

from auditoria_documental.informe_html import generar_informe


def _resultados():
    return [
        {
            "contrato": "20-2026000003",
            "cod": "20",
            "estado_alfresco": "ENCONTRADA",
            "node_id": "n1",
            "encontrados": 1,
            "esperados": 2,
            "faltantes": 1,
            "pct_cumpl": 50.0,
            "filas": [
                {"ID_DOC": "D1", "DOCUMENTO": "Informe de oportunidad", "FORMATO": "FCO.55",
                 "ETAPA": "Precontractual - selección", "ESTADO": "ENCONTRADO",
                 "ARCHIVO": "FCO.55_Informe.pdf"},
                {"ID_DOC": "D2", "DOCUMENTO": "Solicitud de cotizaciones", "FORMATO": "FCO.57",
                 "ETAPA": "Precontractual - selección", "ESTADO": "FALTANTE", "ARCHIVO": ""},
                {"ID_DOC": "D3", "DOCUMENTO": "Acta de finalización", "FORMATO": "FCO.66",
                 "ETAPA": "Cierre y liquidación", "ESTADO": "NO APLICA", "ARCHIVO": ""},
            ],
        },
        {
            "contrato": "270-2026000050",
            "cod": "270",
            "estado_alfresco": "NO SE EVIDENCIA",
            "encontrados": 0,
            "esperados": 1,
            "faltantes": 1,
            "pct_cumpl": 0.0,
            "filas": [
                {"ID_DOC": "D4", "DOCUMENTO": "Acta de liquidación", "FORMATO": "FCO.67",
                 "ETAPA": "Cierre y liquidación", "ESTADO": "FALTANTE", "ARCHIVO": ""},
            ],
        },
    ]


def test_generar_informe_modulos_y_detalle(tmp_path):
    ruta = str(tmp_path / "11_Informe_Auditoria.html")

    generar_informe(_resultados(), ruta, periodo="DEMO", fecha_revision="2026-08-31")

    with open(ruta, encoding="utf-8") as archivo:
        html = archivo.read()

    assert "Auditoría Documental Automatizada" in html
    # Módulos desplegables.
    assert 'details class="modulo"' in html
    assert "Contratos con carpeta en Alfresco" in html
    assert "Contratos sin carpeta en Alfresco" in html
    # Contrato sin carpeta desplegable y en rojo.
    assert 'details class="contrato sin"' in html
    assert "Necesita carpeta" in html
    # Desplegables de documentos.
    assert "Documentos faltantes" in html and "Documentos encontrados" in html
    # Documento hallado con su archivo real y documento faltante.
    assert "FCO.55" in html and "FCO.55_Informe.pdf" in html
    assert "Solicitud de cotizaciones" in html
    # Ambos contratos presentes y enlace a la carpeta.
    assert "20-2026000003" in html and "270-2026000050" in html
    assert "workspace://SpacesStore/n1" in html
    # Tabla general + descarga en Excel y PDF.
    assert 'id="tabla-resumen"' in html
    assert "Descargar Excel" in html
    assert "Descargar PDF" in html and "function imprimir" in html
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet;base64," in html


def test_excel_embebido_es_valido(tmp_path):
    import base64
    import io
    import re

    import openpyxl

    ruta = str(tmp_path / "11_Informe_Auditoria.html")
    generar_informe(_resultados(), ruta, periodo="DEMO")

    with open(ruta, encoding="utf-8") as archivo:
        html = archivo.read()

    coincidencia = re.search(r"base64,([A-Za-z0-9+/=]+)\"", html)
    assert coincidencia, "no se encontró el Excel embebido"

    libro = openpyxl.load_workbook(io.BytesIO(base64.b64decode(coincidencia.group(1))))
    hoja = libro.active
    assert hoja.cell(1, 1).value == "CONTRATO"
    assert hoja.max_row == 3  # encabezado + 2 contratos
