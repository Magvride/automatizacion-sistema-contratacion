# -*- coding: utf-8 -*-
"""Tests del diccionario de documentos (``src/auditoria_documental/diccionario.py``)."""

import openpyxl
import pytest

from auditoria_documental.diccionario import (
    buscar_diccionario,
    cargar_diccionario,
    clase_de_contrato,
    obligatorios_por_clase,
    ordenar_etapas,
)


def _crear_diccionario(tmp_path):
    ruta = tmp_path / "Diccionario_Documentos.xlsx"
    libro = openpyxl.Workbook()

    hoja = libro.active
    hoja.title = "Diccionario"
    hoja.append(
        ["ID_DOC", "CLASE", "CODIGO_CLASE", "DOCUMENTO", "CODIGO_FORMATO",
         "ETAPA", "CARDINALIDAD", "OBLIGATORIEDAD", "ALIAS"]
    )
    hoja.append(["D190", "ORDEN DE COMPRA", "20", "Informe de oportunidad",
                 "FCO.55", "Precontractual - selección", "1", "Obligatorio",
                 "FCO.55 Informe Oportunidad"])
    hoja.append(["D999", "ORDEN DE COMPRA", "20", "Documento opcional",
                 "", "Cierre y liquidación", "1", "Opcional", ""])
    hoja.append(["D027", "OPS PERSONA NATURAL", "18 PN", "Solicitud de cotizaciones",
                 "FCO.57", "Precontractual - selección", "1", "Obligatorio",
                 "FCO.57 Solicitud Cotizaciones"])

    etapas = libro.create_sheet("Matriz_Etapas")
    etapas.append(["ETAPA", "DOCUMENTO", "CÓDIGO_FORMATO", "CARDINALIDAD", "20"])
    etapas.append(["Precontractual - selección", "Informe de oportunidad", "FCO.55", "1", "X"])

    libro.save(ruta)
    return str(ruta)


def test_cargar_diccionario(tmp_path):
    ruta = _crear_diccionario(tmp_path)
    data = cargar_diccionario(ruta)

    assert data["ruta"] == ruta
    assert len(data["diccionario"]) == 3
    primero = data["diccionario"][0]
    assert primero["ID_DOC"] == "D190"
    assert primero["CODIGO_CLASE"] == "20"
    assert primero["CODIGO_FORMATO"] == "FCO.55"


def test_cargar_diccionario_inexistente(tmp_path, monkeypatch):
    import auditoria_documental.diccionario as dic

    monkeypatch.setattr(dic, "MATRIZ_MANUAL_DIR", tmp_path)
    with pytest.raises(FileNotFoundError):
        cargar_diccionario(str(tmp_path / "no_existe.xlsx"))


def test_obligatorios_por_clase(tmp_path):
    data = cargar_diccionario(_crear_diccionario(tmp_path))

    oblig_20 = obligatorios_por_clase(data["diccionario"], "20")
    assert [d["ID_DOC"] for d in oblig_20] == ["D190"]

    oblig_18 = obligatorios_por_clase(data["diccionario"], "18 PN")
    assert [d["ID_DOC"] for d in oblig_18] == ["D027"]


def test_clase_de_contrato():
    assert clase_de_contrato("20-2026000003", "CASA HERMES LTDA") == "20"
    assert clase_de_contrato("270-2026000050", "") == "270"
    assert clase_de_contrato("298-2026000129", "") == "298"
    assert clase_de_contrato("18-2026001182", "NELSON ALEXIS CAYER GIRALDO") == "18 PN"
    assert clase_de_contrato("18-2026001832", "CASA HERMES LTDA") == "18 PJ"
    assert clase_de_contrato("18-2026001832", "ACME S.A.S") == "18 PJ"


def test_clase_de_contrato_con_ceros_a_la_izquierda():
    assert clase_de_contrato("0450-2026000003", "") == "450"
    assert clase_de_contrato("0020-2026000003", "") == "20"
    assert clase_de_contrato("0018-2026001182", "NELSON") == "18 PN"


def test_ordenar_etapas():
    etapas = [
        "Cierre y liquidación",
        "Precontractual - selección",
        "Ejecución y pagos",
    ]
    assert ordenar_etapas(etapas) == [
        "Precontractual - selección",
        "Ejecución y pagos",
        "Cierre y liquidación",
    ]


def _crear_matriz(tmp_path):
    """Workbook con el formato de la matriz entregada (CATALOGO_PLANO/MATRIZ)."""
    ruta = tmp_path / "Matriz_Documentos_por_Clase.xlsx"
    libro = openpyxl.Workbook()

    hoja = libro.active
    hoja.title = "CATALOGO_PLANO"
    hoja.append(["ID_DOC", "CLASE", "CODIGO_CLASE", "DOCUMENTO", "CODIGO_FORMATO",
                 "ETAPA", "CARDINALIDAD", "OBLIGATORIEDAD", "ALIAS"])
    hoja.append(["D001", "ORDEN DE COMPRA", "20", "Informe de oportunidad",
                 "FCO.55", "Precontractual - selección", "Única", "Obligatorio",
                 "FCO.55 Informe Oportunidad"])
    hoja.append(["D003", "ORDEN DE COMPRA", "20", "Cotizaciones", "",
                 "Precontractual - selección", "Una o varias", "Condicional",
                 "Cotizacion"])
    hoja.append(["D002", "OPS NATURAL", "18 PN", "Solicitud de cotizaciones",
                 "FCO.57", "Precontractual - selección", "Única", "Obligatorio",
                 "FCO.57 Solicitud Cotizaciones"])

    matriz = libro.create_sheet("MATRIZ")
    matriz.append(["ETAPA", "DOCUMENTO", "CÓDIGO FORMATO", "CARDINALIDAD",
                   "298\nSuministro", "20\nOrden compra"])
    matriz.append(["Precontractual - selección", "Informe de oportunidad",
                   "FCO.55", "Única", "X", "X"])

    libro.save(ruta)
    return str(ruta)


def test_cargar_diccionario_formato_matriz(tmp_path):
    data = cargar_diccionario(_crear_matriz(tmp_path))

    assert len(data["diccionario"]) == 3
    assert [d["ID_DOC"] for d in obligatorios_por_clase(data["diccionario"], "20")] == ["D001"]
    assert [d["ID_DOC"] for d in obligatorios_por_clase(data["diccionario"], "18 PN")] == ["D002"]


def test_buscar_diccionario_explicito(tmp_path):
    ruta = _crear_matriz(tmp_path)
    assert buscar_diccionario(ruta) == ruta


def test_buscar_diccionario_inexistente(tmp_path, monkeypatch):
    import auditoria_documental.diccionario as dic

    monkeypatch.setattr(dic, "MATRIZ_MANUAL_DIR", tmp_path)
    assert buscar_diccionario(str(tmp_path / "nada.xlsx")) == ""


def test_cargar_diccionario_bloqueado(tmp_path, monkeypatch):
    import auditoria_documental.diccionario as dic

    ruta = tmp_path / "Diccionario_Documentos.xlsx"
    ruta.write_bytes(b"contenido")

    def _bloqueado(*args, **kwargs):
        raise PermissionError("bloqueado por Excel")

    monkeypatch.setattr(dic.openpyxl, "load_workbook", _bloqueado)
    monkeypatch.setattr("time.sleep", lambda *_: None)

    with pytest.raises(PermissionError) as excinfo:
        dic.cargar_diccionario(str(ruta))
    assert "abierto o bloqueado" in str(excinfo.value)

