# -*- coding: utf-8 -*-
"""Tests de los borradores de correo (``src/auditoria_documental/correos.py``)."""

import openpyxl

from auditoria_documental.correos import (
    cargar_plantillas,
    construir_mensajes,
    escribir_borradores,
    escribir_borradores_excel,
    extraer_correos_apoyo,
    renderizar,
)


def _workbook_correos(tmp_path):
    ruta = tmp_path / "Matriz_Documentos_por_Clase.xlsx"
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "CORREOS"
    filas = [
        ["PROFORMAS DE CORREO", None],
        ["Campos entre llaves dobles", None],
        [None, None],
        ["CAMPOS DE COMBINACIÓN", None],
        ["{{CONTRATO}}", "Columna CONTRATO"],
        [None, None],
        ["1. SOLICITUD DE DOCUMENTACIÓN FALTANTE", None],
        ["Cuándo se envía", "Cuando el ESTADO es Incompleto."],
        ["Para", "Supervisor del contrato ({{SUPERVISOR}})."],
        ["Asunto", "Solicitud de documentación faltante. Contrato {{CONTRATO}}"],
        ["Cuerpo", "Cordial saludo, {{SUPERVISOR}}.\nFaltantes: {{LISTADO_FALTANTES}}. Completitud: {{PORCENTAJE}}. Enlace: {{ENLACE_CARPETA}}."],
        [None, None],
        ["3. CARPETA NO LOCALIZADA O NO DESCARGADA", None],
        ["Cuándo se envía", "Cuando la carpeta no se encuentra."],
        ["Para", "Auxiliar del equipo."],
        ["Asunto", "Carpeta no localizada en Alfresco. Contrato {{CONTRATO}}"],
        ["Cuerpo", "No fue posible localizar la carpeta del contrato {{CONTRATO}} ({{CLASE}})."],
        [None, None],
        ["4. EXPEDIENTE COMPLETO Y CONFORME", None],
        ["Cuándo se envía", "Cuando el ESTADO es Completo."],
        ["Para", "Supervisor del contrato."],
        ["Asunto", "Expediente completo. Contrato {{CONTRATO}}"],
        ["Cuerpo", "El expediente del contrato {{CONTRATO}} está completo."],
    ]
    for fila in filas:
        hoja.append(fila)
    libro.save(ruta)
    return str(ruta)


def _contrato_incompleto():
    return {
        "contrato": "20-2026000003",
        "cod": "20",
        "estado_alfresco": "ENCONTRADA",
        "node_id": "n1",
        "supervisor": "JUAN PEREZ",
        "correo_ordenador": "juan@uis.edu.co",
        "filas": [
            {"ID_DOC": "D190", "DOCUMENTO": "Informe oportunidad", "FORMATO": "FCO.55",
             "ESTADO": "ENCONTRADO", "ARCHIVO": "FCO.55.pdf"},
            {"ID_DOC": "D027", "DOCUMENTO": "Solicitud cotizaciones", "FORMATO": "FCO.57",
             "ESTADO": "FALTANTE", "ARCHIVO": ""},
        ],
    }


def _contrato_no_encontrado():
    return {
        "contrato": "20-2026009999",
        "cod": "20",
        "estado_alfresco": "NO SE EVIDENCIA",
        "filas": [],
    }


def test_renderizar():
    assert renderizar("Hola {{SUPERVISOR}}, {{NO_EXISTE}}!", {"SUPERVISOR": "JUAN"}) == "Hola JUAN, !"


def test_cargar_plantillas(tmp_path):
    plantillas = cargar_plantillas(_workbook_correos(tmp_path))
    assert len(plantillas) == 3
    assert plantillas[0]["titulo"] == "SOLICITUD DE DOCUMENTACIÓN FALTANTE"
    assert plantillas[0]["asunto"] == "Solicitud de documentación faltante. Contrato {{CONTRATO}}"
    assert "{{LISTADO_FALTANTES}}" in plantillas[0]["cuerpo"]


def test_construir_mensajes_selecciona_plantilla(tmp_path):
    plantillas = cargar_plantillas(_workbook_correos(tmp_path))
    mensajes = construir_mensajes(
        [_contrato_incompleto(), _contrato_no_encontrado()],
        plantillas,
        fecha_revision="2026-08-31",
    )
    assert len(mensajes) == 2

    faltante = mensajes[0]
    assert faltante["tipo"] == "SOLICITUD DE DOCUMENTACIÓN FALTANTE"
    assert "20-2026000003" in faltante["asunto"]
    assert "JUAN PEREZ" in faltante["cuerpo"]
    assert "FCO.57 Solicitud cotizaciones" in faltante["cuerpo"]
    assert "50.0%" in faltante["cuerpo"]

    no_loc = mensajes[1]
    assert no_loc["tipo"] == "CARPETA NO LOCALIZADA O NO DESCARGADA"


def test_escribir_borradores(tmp_path):
    plantillas = cargar_plantillas(_workbook_correos(tmp_path))
    mensajes = construir_mensajes([_contrato_incompleto()], plantillas)
    ruta = str(tmp_path / "10_Correos_Auditoria.txt")

    escribir_borradores(mensajes, ruta)

    with open(ruta, encoding="utf-8") as archivo:
        contenido = archivo.read()
    assert "PARA: JUAN PEREZ" in contenido
    assert "ASUNTO: Documentación pendiente del contrato 20-2026000003" in contenido


def test_escribir_borradores_excel_para_power_automate(tmp_path):
    mensajes = [{
        "para": "JUAN PEREZ",
        "contrato": "20-2026000003",
        "correo": "juan@uis.edu.co",
        "correos_apoyo": "apoyo1@uis.edu.co; apoyo2@uis.edu.co",
        "tipo": "CARPETA NO LOCALIZADA O NO DESCARGADA",
        "centro_costo": "CC1",
        "ordenador_centro_costo": "JUAN PEREZ - CC1",
        "contratista": "ACME S.A.S.",
        "fecha_contrato": "2026/01/15",
        "mes_documento": "agosto",
        "mes_numero": "08",
        "dias_transcurridos": 243,
        "objeto": "Prestación de servicios profesionales",
        "asunto": "Cordial saludo, JUAN PEREZ. Documentación pendiente del contrato 20-2026000003",
        "cuerpo": "Cordial saludo. Faltan documentos del contrato 20-2026000003.",
    }]
    ruta = str(tmp_path / "10_Correos_Auditoria.xlsx")

    escribir_borradores_excel(mensajes, ruta)

    libro = openpyxl.load_workbook(ruta)
    hoja = libro["Correos"]
    assert [celda.value for celda in hoja[1]] == ["CORREO", "CC", "ASUNTO", "CUERPO"]
    assert hoja.cell(row=2, column=1).value == "juan@uis.edu.co"
    assert hoja.cell(row=2, column=2).value == "apoyo1@uis.edu.co; apoyo2@uis.edu.co"
    assert hoja.cell(row=2, column=3).value == (
        "Documentación pendiente del contrato 20-2026000003"
    )
    assert "Cordial saludo" not in hoja.cell(row=2, column=3).value
    assert "JUAN PEREZ" not in hoja.cell(row=2, column=3).value
    assert "En el marco del seguimiento contractual del mes de agosto" in hoja.cell(row=2, column=4).value
    assert hoja.cell(row=2, column=4).value.startswith("<p>Cordial saludo, JUAN PEREZ.</p>")
    assert "<ol>" in hoja.cell(row=2, column=4).value
    assert "<li><strong>1. Contrato 20-2026000003</strong>" in hoja.cell(row=2, column=4).value
    assert "CC1" not in hoja.cell(row=2, column=4).value
    assert "ACME S.A.S." in hoja.cell(row=2, column=4).value
    assert "2026/01/15" in hoja.cell(row=2, column=4).value
    assert "Prestación de servicios profesionales" in hoja.cell(row=2, column=4).value
    assert "20-2026000003" in hoja.cell(row=2, column=4).value
    assert "aún no ha sido cargado" in hoja.cell(row=2, column=4).value
    assert hoja.tables["CorreosPowerAutomate"].ref == "A1:D2"


def test_escribir_borradores_excel_conserva_contratos_del_mismo_correo(tmp_path):
    mensajes = [
        {
            "para": "anaberam@uis.edu.co",
            "correo": "anaberam@uis.edu.co",
            "contrato": "20-2026000436",
            "tipo": "CARPETA NO LOCALIZADA O NO DESCARGADA",
            "contratista": "LACTEOS ROVIRENSES S.A",
            "fecha_contrato": "2026/09/14",
            "mes_documento": "septiembre",
            "mes_numero": "09",
            "dias_transcurridos": 1,
            "objeto": "COMPRA DE HARINA DE MAÍZ",
        },
        {
            "para": "anaberam@uis.edu.co",
            "correo": "anaberam@uis.edu.co",
            "contrato": "20-2026000437",
            "tipo": "CARPETA NO LOCALIZADA O NO DESCARGADA",
            "contratista": "OTRO CONTRATISTA",
            "fecha_contrato": "2026/09/13",
            "mes_documento": "septiembre",
            "mes_numero": "09",
            "dias_transcurridos": 2,
            "objeto": "OTRO OBJETO",
        },
    ]
    ruta = str(tmp_path / "10_Correos_Auditoria.xlsx")

    escribir_borradores_excel(mensajes, ruta)

    libro = openpyxl.load_workbook(ruta)
    hoja = libro["Correos"]
    assert hoja.max_row == 2
    assert hoja.cell(row=2, column=1).value == "anaberam@uis.edu.co"
    assert hoja.cell(row=2, column=3).value == "Documentación pendiente en la verificación de alfresco"
    assert "Cordial saludo" not in hoja.cell(row=2, column=3).value
    cuerpo = hoja.cell(row=2, column=4).value
    assert "20-2026000436" in cuerpo
    assert "20-2026000437" in cuerpo
    assert cuerpo.count("Cordial saludo,") == 1
    assert cuerpo.count("En el marco del seguimiento contractual") == 1
    assert cuerpo.count("Gracias por su colaboración.") == 1
    assert cuerpo.count("Le agradecemos gestionar el cargue") == 1
    assert "<ol>" in cuerpo and "<li>" in cuerpo
    assert hoja.tables["CorreosPowerAutomate"].ref == "A1:D2"


def test_cuerpo_html_escapa_valores(tmp_path):
    mensajes = [{
        "para": "JUAN <jefe> PEREZ & CIA",
        "contrato": "20-1",
        "correo": "juan@uis.edu.co",
        "tipo": "CARPETA NO LOCALIZADA O NO DESCARGADA",
        "contratista": "ACME <S.A.S.>",
        "fecha_contrato": "2026/01/15",
        "mes_documento": "agosto",
        "mes_numero": "08",
        "dias_transcurridos": 1,
        "objeto": "Servicios & consultoría",
    }]
    ruta = str(tmp_path / "10_Correos_Auditoria.xlsx")

    escribir_borradores_excel(mensajes, ruta)

    cuerpo = openpyxl.load_workbook(ruta)["Correos"].cell(row=2, column=4).value
    assert "JUAN <jefe>" not in cuerpo
    assert "JUAN &lt;jefe&gt; PEREZ &amp; CIA" in cuerpo
    assert "ACME &lt;S.A.S.&gt;" in cuerpo


def test_extraer_correos_apoyo_formatos_sep():
    assert extraer_correos_apoyo("apoyo1@uis.edu.co; apoyo2@uis.edu.co") == [
        "apoyo1@uis.edu.co", "apoyo2@uis.edu.co",
    ]
    assert extraer_correos_apoyo("a@uis.edu.co\nb@uis.edu.co\nc@uis.edu.co") == [
        "a@uis.edu.co", "b@uis.edu.co", "c@uis.edu.co",
    ]
    assert extraer_correos_apoyo("ESCUELA DE FILOSOFIA <escfilosofia@uis.edu.co>") == [
        "escfilosofia@uis.edu.co",
    ]
    assert extraer_correos_apoyo("barinf@uis.edu.co>\n <bartesor@uis.edu.co>") == [
        "barinf@uis.edu.co", "bartesor@uis.edu.co",
    ]
    # Sin duplicados aunque se repita el correo.
    assert extraer_correos_apoyo("a@uis.edu.co, A@uis.edu.co") == ["a@uis.edu.co"]
    assert extraer_correos_apoyo(None) == []
    assert extraer_correos_apoyo("sin correos aquí") == []


def test_escribir_borradores_excel_une_apoyos_del_grupo(tmp_path):
    mensajes = [
        {
            "para": "JUAN PEREZ", "contrato": "20-1", "correo": "juan@uis.edu.co",
            "correos_apoyo": "apoyo1@uis.edu.co",
            "tipo": "CARPETA NO LOCALIZADA O NO DESCARGADA",
            "contratista": "A", "fecha_contrato": "2026/01/15",
            "mes_documento": "agosto", "mes_numero": "08",
            "dias_transcurridos": 1, "objeto": "O1",
        },
        {
            "para": "JUAN PEREZ", "contrato": "20-2", "correo": "juan@uis.edu.co",
            "correos_apoyo": "apoyo2@uis.edu.co; apoyo1@uis.edu.co; juan@uis.edu.co",
            "tipo": "CARPETA NO LOCALIZADA O NO DESCARGADA",
            "contratista": "B", "fecha_contrato": "2026/01/16",
            "mes_documento": "agosto", "mes_numero": "08",
            "dias_transcurridos": 2, "objeto": "O2",
        },
    ]
    ruta = str(tmp_path / "10_Correos_Auditoria.xlsx")

    escribir_borradores_excel(mensajes, ruta)

    libro = openpyxl.load_workbook(ruta)
    hoja = libro["Correos"]
    assert hoja.max_row == 2
    # Unión sin duplicados y sin repetir el correo principal.
    assert hoja.cell(row=2, column=2).value == "apoyo1@uis.edu.co; apoyo2@uis.edu.co"
