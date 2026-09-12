# -*- coding: utf-8 -*-
"""
Genera ``Diccionario_Documentos.xlsx`` (canónico) a partir de la matriz de
documentos por clase.

Produce un archivo con las hojas que el motor espera:

    - Diccionario    (catálogo: una fila por documento y clase)
    - Matriz_Etapas  (matriz etapa × clase)
    - CORREOS        (plantillas de correo)
    - REGLAS         (reglas de negocio del matching, documentadas)

Uso:
    python scripts/generar_diccionario.py
    python scripts/generar_diccionario.py <origen.xlsx> <destino.xlsx>
"""

import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "src"))

from config import MATRIZ_MANUAL_DIR  # noqa: E402

ORIGEN_DEFECTO = MATRIZ_MANUAL_DIR / "Matriz_Documentos_por_Clase.xlsx"
DESTINO_DEFECTO = MATRIZ_MANUAL_DIR / "Diccionario_Documentos.xlsx"

# Hoja de origen -> hoja de destino en el diccionario canónico.
MAPEO_HOJAS = {
    "CATALOGO_PLANO": "Diccionario",
    "Diccionario": "Diccionario",
    "MATRIZ": "Matriz_Etapas",
    "Matriz_Etapas": "Matriz_Etapas",
    "CORREOS": "CORREOS",
}

FUENTE_HEADER = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
RELLENO_HEADER = PatternFill("solid", fgColor="1F4E78")
FUENTE_TITULO = Font(name="Calibri", size=13, bold=True, color="1F4E78")
FUENTE_DATO = Font(name="Calibri", size=10)

REGLAS = [
    ("DICCIONARIO DE DOCUMENTOS — REGLAS DE NEGOCIO", ""),
    ("", ""),
    ("Qué es", "Lista de documentos exigibles por tipo de contrato. El motor de "
               "auditoría compara los archivos reales de cada carpeta de Alfresco "
               "contra estas reglas y marca ENCONTRADO / FALTANTE / NO APLICA."),
    ("Hoja principal", "'Diccionario' (una fila por documento y clase)."),
    ("Hoja de etapas", "'Matriz_Etapas' (marca con X qué documento aplica a cada clase)."),
    ("Hoja de correos", "'CORREOS' (plantillas con campos {{CONTRATO}}, {{SUPERVISOR}}, "
                        "{{LISTADO_FALTANTES}}, {{PORCENTAJE}}, {{ENLACE_CARPETA}}, ...)."),
    ("", ""),
    ("Columnas del catálogo", ""),
    ("ID_DOC", "Identificador del documento (ej. D190)."),
    ("CLASE", "Nombre legible de la clase (ej. ORDEN DE COMPRA)."),
    ("CODIGO_CLASE", "Clase del contrato a la que aplica: 20, 19, 298, 450, 270, "
                     "18 PN (persona natural), 18 PJ (persona jurídica)."),
    ("DOCUMENTO", "Nombre del documento exigido."),
    ("CODIGO_FORMATO", "Código de formato FCO/FTH/FFI (ej. FCO.55). Es la coincidencia "
                       "más fuerte contra el nombre del archivo."),
    ("ETAPA", "Etapa del proceso: Precontractual - selección / idoneidad / presupuestal, "
              "Contractual / perfeccionamiento, Ejecución y pagos, Cierre y liquidación."),
    ("CARDINALIDAD", "Única, Una o varias, Condicional, Única + por pago, etc."),
    ("OBLIGATORIEDAD", "Obligatorio (se exige) o Condicional (no cuenta como faltante)."),
    ("ALIAS", "Nombres alternativos separados por coma; se usan para reconocer archivos "
              "cuyo nombre no trae el código de formato."),
    ("", ""),
    ("Reglas de coincidencia (matching)", ""),
    ("1. Código de formato", "Si el documento tiene FCO/FTH/FFI, se exige el código exacto "
                             "(ej. fco55) en el nombre del archivo."),
    ("2. Alias completo", "Si no está el código, se acepta que el alias completo (≥2 tokens "
                          "distintivos) aparezca en el nombre del archivo."),
    ("3. Sin formato", "Se exigen los tokens distintivos del documento/alias (≥2 si la frase "
                       "es multi-token); se ignoran acentos, signos y números como (32)."),
    ("4. Equivalencias", "Seguridad Social ↔ Parafiscales; ARL ↔ Estándares/Constancia. "
                         "Se evalúan por tokens completos (no por subcadena)."),
    ("5. Regla de cierre", "Si en un contrato tipo 18 existe FCO.74 (acta de pago final), "
                           "FCO.66 y FCO.67 quedan NO APLICA (no son faltantes)."),
    ("6. Clase 18", "PN vs PJ se infiere del contratista (SAS/LTDA/S.A/S.A.S = persona jurídica)."),
    ("", ""),
    ("Mínimo viable", "Una hoja 'Diccionario' con CODIGO_CLASE, DOCUMENTO y OBLIGATORIEDAD. "
                      "Con CODIGO_FORMATO, ETAPA y ALIAS el resultado es mucho más preciso."),
]


def _copiar_hoja(libro_origen, hoja_origen, libro_destino, nombre_destino):
    hoja = libro_origen[hoja_origen]
    destino = libro_destino.create_sheet(nombre_destino)

    filas = list(hoja.iter_rows(values_only=True))
    if not filas:
        return

    for fila in filas:
        destino.append(list(fila))

    for col in range(1, destino.max_column + 1):
        celda = destino.cell(row=1, column=col)
        celda.font = FUENTE_HEADER
        celda.fill = RELLENO_HEADER
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col in range(1, destino.max_column + 1):
        ancho = 12
        for fila in range(1, min(destino.max_row, 200) + 1):
            valor = destino.cell(row=fila, column=col).value
            if valor is not None:
                ancho = max(ancho, min(len(str(valor)) + 2, 60))
        destino.column_dimensions[get_column_letter(col)].width = ancho

    destino.freeze_panes = "A2"
    destino.auto_filter.ref = destino.dimensions


def _escribir_reglas(libro):
    hoja = libro.create_sheet("REGLAS")
    hoja.column_dimensions["A"].width = 30
    hoja.column_dimensions["B"].width = 110

    for i, (titulo, texto) in enumerate(REGLAS, start=1):
        celda_titulo = hoja.cell(row=i, column=1, value=titulo)
        celda_texto = hoja.cell(row=i, column=2, value=texto)
        if texto == "" and titulo:
            celda_titulo.font = FUENTE_TITULO
        else:
            celda_titulo.font = Font(name="Calibri", size=10, bold=True)
            celda_texto.font = FUENTE_DATO
        celda_texto.alignment = Alignment(wrap_text=True, vertical="top")


def main() -> int:
    origen = Path(sys.argv[1]) if len(sys.argv) > 1 else ORIGEN_DEFECTO
    destino = Path(sys.argv[2]) if len(sys.argv) > 2 else DESTINO_DEFECTO

    if not origen.is_file():
        print(f"[ERROR] No existe el archivo de origen: {origen}")
        return 1

    libro_origen = openpyxl.load_workbook(origen, data_only=True, read_only=True)
    try:
        libro_destino = openpyxl.Workbook()
        libro_destino.remove(libro_destino.active)

        copiadas = []
        for hoja_origen, nombre_destino in MAPEO_HOJAS.items():
            if hoja_origen in libro_origen.sheetnames and nombre_destino not in copiadas:
                _copiar_hoja(libro_origen, hoja_origen, libro_destino, nombre_destino)
                copiadas.append(nombre_destino)

        _escribir_reglas(libro_destino)
    finally:
        libro_origen.close()

    destino.parent.mkdir(parents=True, exist_ok=True)
    libro_destino.save(destino)

    print(f"[OK] Diccionario generado: {destino}")
    print(f"     Hojas: {', '.join(libro_destino.sheetnames)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
