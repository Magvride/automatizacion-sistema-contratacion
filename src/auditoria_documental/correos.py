# -*- coding: utf-8 -*-
"""
Borradores de correo a partir del diagnóstico documental.

Carga las plantillas de la hoja ``CORREOS`` del Excel de documentos y las
combina con las columnas de la hoja ``DIAGNOSTICO`` (ver
``auditoria_documental.workbook.fila_diagnostico``).

Plantillas reconocidas por su título: solicitud de faltantes, inconsistencia con
la matriz, carpeta no localizada, expediente completo y consolidado por unidad.
La selección por contrato depende de ``ESTADO`` y de
``COINCIDE CON MATRIZ DE SEGUIMIENTO``.
"""

import os
import re
import unicodedata
from datetime import date, datetime
from html import escape as _escape_html

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from utils.logger import configurar_logger
from .diccionario import _abrir_libro, buscar_diccionario
from .workbook import fila_diagnostico

logger = configurar_logger("correos")

HOJAS_CORREOS = ("CORREOS",)

# Tipos de plantilla (se detectan por palabra clave en el título).
TIPOS = ("faltante", "inconsistencia", "no_localizada", "completo", "consolidado")

_CAMPO_RE = re.compile(r"\{\{\s*([A-Z_]+)\s*\}\}")

# Mismo patrón que ``ordenadores.extraer_correos``: las celdas de apoyo del
# reporte SEP mezclan nombres y formatos ("ESCUELA X <apoyo@uis.edu.co>").
_CORREO_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")


def extraer_correos_apoyo(texto) -> list:
    """Extrae los correos de apoyo de una celda libre, sin duplicados y en orden."""
    if texto is None:
        return []
    vistos = set()
    correos = []
    for encontrado in _CORREO_RE.findall(str(texto)):
        correo = encontrado.strip()
        clave = correo.casefold()
        if clave and clave not in vistos:
            vistos.add(clave)
            correos.append(correo)
    return correos


def asunto_correo(contrato: str = "", cantidad: int = 1) -> str:
    """Asunto del correo: sin saludo ni nombre del ordenador."""
    if cantidad > 1:
        return "Documentación pendiente en la verificación de alfresco"
    return f"Documentación pendiente del contrato {str(contrato or '').strip()}".strip()


def _clave(texto) -> str:
    """Mayúsculas sin acentos para comparar etiquetas."""
    t = unicodedata.normalize("NFKD", str(texto or ""))
    t = "".join(c for c in t if not unicodedata.combining(c))
    return t.strip().upper()


def _parsear(hoja) -> list:
    """Convierte la hoja CORREOS en una lista de plantillas."""
    plantillas = []
    actual = None
    for fila in hoja.iter_rows(values_only=True):
        c1 = str(fila[0]).strip() if fila and fila[0] is not None else ""
        c2 = str(fila[1]).strip() if fila and len(fila) > 1 and fila[1] is not None else ""

        if re.match(r"^\d+\.\s", c1):
            actual = {
                "titulo": re.sub(r"^\d+\.\s*", "", c1),
                "cuando": "",
                "para": "",
                "asunto": "",
                "cuerpo": "",
            }
            plantillas.append(actual)
            continue
        if actual is None:
            continue

        etiqueta = _clave(c1)
        if etiqueta == "CUANDO SE ENVIA":
            actual["cuando"] = c2
        elif etiqueta == "PARA":
            actual["para"] = c2
        elif etiqueta == "ASUNTO":
            actual["asunto"] = c2
        elif etiqueta == "CUERPO":
            actual["cuerpo"] = c2

    return [p for p in plantillas if p["cuerpo"]]


def cargar_plantillas(ruta: str = "") -> list:
    """Carga las plantillas de la hoja CORREOS (autodetecta el archivo si vacío)."""
    ruta = buscar_diccionario(ruta)
    if not ruta:
        logger.warning("No se encontró el archivo de documentos para leer las plantillas de correo.")
        return []

    libro = _abrir_libro(ruta)
    try:
        for nombre in HOJAS_CORREOS:
            if nombre in libro.sheetnames:
                plantillas = _parsear(libro[nombre])
                logger.info("Plantillas de correo cargadas: %d.", len(plantillas))
                return plantillas
    finally:
        libro.close()
    logger.warning("El archivo no contiene la hoja %s.", " / ".join(HOJAS_CORREOS))
    return []


def renderizar(texto, valores: dict) -> str:
    """Reemplaza los campos ``{{CAMPO}}`` por su valor (vacío si no existe)."""
    return _CAMPO_RE.sub(lambda m: str(valores.get(m.group(1), "")), str(texto or ""))


def _tipo_de(plantilla: dict) -> str:
    titulo = plantilla.get("titulo", "").lower()
    for tipo in TIPOS:
        palabra = tipo.replace("_", " ")
        if palabra in titulo:
            return tipo
    return ""


def _por_tipo(plantillas: list, tipo: str):
    for plantilla in plantillas:
        if _tipo_de(plantilla) == tipo:
            return plantilla
    return None


def elegir_plantilla(plantillas: list, fila: dict):
    """Selecciona la plantilla adecuada según el estado del contrato."""
    estado = str(fila.get("ESTADO", "")).strip()
    carpeta = str(fila.get("CARPETA ENCONTRADA", "")).strip().lower()
    coincide = str(fila.get("COINCIDE CON MATRIZ DE SEGUIMIENTO", "")).strip().lower()

    if carpeta.startswith("no encontrada") or estado == "No encontrada":
        return _por_tipo(plantillas, "no_localizada")
    if coincide.startswith("no coincide"):
        return _por_tipo(plantillas, "inconsistencia")

    faltantes = str(fila.get("DOCS FALTANTES", "")).strip()
    try:
        tiene_faltantes = int(float(faltantes or 0)) > 0
    except (TypeError, ValueError):
        tiene_faltantes = bool(faltantes and faltantes.lower() not in ("no evaluado", "0"))

    if estado == "Completo" or (carpeta.startswith("encontrada") and not tiene_faltantes):
        return _por_tipo(plantillas, "completo")
    if estado == "Incompleto" or tiene_faltantes:
        return _por_tipo(plantillas, "faltante")
    return None


def construir_mensajes(
    contratos: list,
    plantillas: list,
    fecha_revision: str = "",
    fecha_limite: str = "",
    unidad_por_contrato: dict = None,
) -> list:
    """Construye los borradores de correo (uno por contrato que lo requiera)."""
    unidad_por_contrato = unidad_por_contrato or {}
    mensajes = []
    for contrato in contratos or []:
        fila = fila_diagnostico(contrato, fecha_revision)
        plantilla = elegir_plantilla(plantillas, fila)
        if not plantilla:
            continue

        contrato_id = fila["CONTRATO"]
        valores = {
            "CONTRATO": contrato_id,
            "CLASE": fila["CLASE"],
            "SUPERVISOR": fila["SUPERVISOR / DESTINATARIO"],
            "ENLACE_CARPETA": fila["ENLACE CARPETA ALFRESCO"],
            "LISTADO_FALTANTES": fila["LISTADO DE FALTANTES"],
            "PORCENTAJE": fila["% COMPLETITUD"],
            "OBSERVACION": fila["OBSERVACIÓN (PARA EL CORREO)"],
            "FECHA_LIMITE": fecha_limite,
            "UNIDAD": unidad_por_contrato.get(contrato_id, contrato.get("centro_costo", "")),
        }
        datos_fecha = _datos_fecha(contrato.get("fecha_contrato", ""), fecha_revision)
        mensajes.append(
            {
                "contrato": contrato_id,
                "tipo": plantilla["titulo"],
                "para": fila["SUPERVISOR / DESTINATARIO"],
                "correo": fila["CORREO ORDENADOR"],
                # Correos de apoyo del reporte SEP: van como CC en el Excel
                # de envío y también son destinatarios del aviso.
                "correos_apoyo": contrato.get("correos_apoyo", ""),
                "asunto": asunto_correo(contrato_id),
                "cuerpo": renderizar(plantilla["cuerpo"], valores),
                "ordenador_centro_costo": " - ".join(
                    parte for parte in (
                        fila["SUPERVISOR / DESTINATARIO"],
                        fila["CENTRO DE COSTO"],
                    ) if str(parte).strip()
                ),
                "centro_costo": fila["CENTRO DE COSTO"],
                "contratista": contrato.get("contratista", ""),
                "objeto": contrato.get("objeto", ""),
                "listado_faltantes": fila["LISTADO DE FALTANTES"],
                **datos_fecha,
            }
        )
    return mensajes


def escribir_borradores(mensajes: list, ruta: str) -> str:
    """Escribe los borradores en un archivo de texto."""
    lineas = []
    for mensaje in mensajes:
        lineas.append("=" * 72)
        lineas.append(f"PARA: {mensaje['para']}")
        lineas.append(f"ASUNTO: {asunto_correo(mensaje.get('contrato', ''))}")
        lineas.append("")
        lineas.append(mensaje["cuerpo"])
        lineas.append("")

    os.makedirs(os.path.dirname(ruta) or ".", exist_ok=True)
    with open(ruta, "w", encoding="utf-8") as archivo:
        archivo.write("\n".join(lineas))
    logger.info("Borradores de correo escritos: %s (%d).", ruta, len(mensajes))
    return ruta


def escribir_borradores_excel(mensajes: list, ruta: str) -> str:
    """Escribe los borradores en una tabla compatible con Power Automate.

    Columnas: ``CORREO`` (ordenador), ``CC`` (correos de apoyo del reporte
    SEP, también destinatarios), ``ASUNTO`` y ``CUERPO``.

    El guardado es atómico (temporal + reemplazo) y, si el destino está
    bloqueado (Excel abierto o sincronización de OneDrive -> ``PermissionError``),
    guarda una copia versionada ``..._YYYYMMDD_HHMMSS.xlsx`` en la misma carpeta
    y la devuelve, para no dejar el archivo viejo como si fuera el nuevo.
    """
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Correos"
    cabeceras = ["CORREO", "CC", "ASUNTO", "CUERPO"]
    hoja.append(cabeceras)

    filas = _agrupar_mensajes_excel(mensajes)
    for correo, cc, asunto, cuerpo in filas:
        hoja.append([
            correo,
            cc,
            asunto,
            cuerpo,
        ])

    encabezado = PatternFill("solid", fgColor="1F4E78")
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = encabezado
        celda.alignment = Alignment(horizontal="center", vertical="center")

    hoja.column_dimensions["A"].width = 34
    hoja.column_dimensions["B"].width = 40
    hoja.column_dimensions["C"].width = 60
    hoja.column_dimensions["D"].width = 110
    hoja.freeze_panes = "A2"

    if filas:
        ultima_fila = len(filas) + 1
        tabla = Table(displayName="CorreosPowerAutomate", ref=f"A1:D{ultima_fila}")
        tabla.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        hoja.add_table(tabla)

    for fila in hoja.iter_rows(min_row=2, max_row=hoja.max_row, min_col=1, max_col=4):
        for celda in fila:
            celda.alignment = Alignment(vertical="top", wrap_text=True)

    os.makedirs(os.path.dirname(ruta) or ".", exist_ok=True)
    try:
        tmp = f"{ruta}.tmp"
        libro.save(tmp)
        os.replace(tmp, ruta)
    except PermissionError as exc:
        from datetime import datetime as _dt

        base, ext = os.path.splitext(ruta)
        alternativa = f"{base}_{_dt.now().strftime('%Y%m%d_%H%M%S')}{ext or '.xlsx'}"
        libro.save(alternativa)
        logger.warning(
            "No se pudo sobrescribir %s (¿abierto en Excel/OneDrive?): %s. "
            "Última versión guardada en %s. Cierre el archivo y regenere.",
            ruta, exc, alternativa,
        )
        return alternativa
    logger.info("Excel de correos para Power Automate guardado: %s (%d).", ruta, len(filas))
    return ruta


_CIERRE_CORREO = (
    "Le agradecemos gestionar el cargue de la documentación a la brevedad posible, "
    "con el fin de mantener los expedientes al día y evitar retrasos en la rendición. "
    "Cualquier novedad que impida el cargue del contrato, le agradecemos notificarlo.\n\n"
    "Gracias por su colaboración."
)

# Cierre en dos párrafos HTML (mismo texto que ``_CIERRE_CORREO``).
_CIERRE_CORREO_HTML_1 = (
    "Le agradecemos gestionar el cargue de la documentación a la brevedad posible, "
    "con el fin de mantener los expedientes al día y evitar retrasos en la rendición. "
    "Cualquier novedad que impida el cargue del contrato, le agradecemos notificarlo."
)
_CIERRE_CORREO_HTML_2 = "Gracias por su colaboración."


def _agrupar_mensajes_excel(mensajes: list) -> list:
    """Agrupa en un solo correo los contratos del mismo destinatario.

    Devuelve filas ``(correo, cc, asunto, cuerpo)`` donde ``cc`` une los
    correos de apoyo SEP de todos los contratos del grupo (sin duplicados y
    sin repetir el correo principal). Los apoyos también son destinatarios.
    """
    grupos = {}
    orden = []
    for indice, mensaje in enumerate(mensajes or []):
        correo = str(mensaje.get("correo", "")).strip()
        clave = correo.casefold() or f"__sin_correo_{indice}"
        if clave not in grupos:
            grupos[clave] = {"correo": correo, "mensajes": [mensaje], "apoyos": []}
            orden.append(clave)
        else:
            grupos[clave]["mensajes"].append(mensaje)
        for apoyo in extraer_correos_apoyo(mensaje.get("correos_apoyo", "")):
            ya = {a.casefold() for a in grupos[clave]["apoyos"]}
            if apoyo.casefold() not in ya and apoyo.casefold() != correo.casefold():
                grupos[clave]["apoyos"].append(apoyo)

    filas = []
    for clave in orden:
        grupo = grupos[clave]
        lote = grupo["mensajes"]
        contrato = str(lote[0].get("contrato", "")).strip()
        asunto = asunto_correo(contrato=contrato, cantidad=len(lote))
        filas.append((grupo["correo"], "; ".join(grupo["apoyos"]), asunto, _cuerpo_power_automate(lote)))
    return filas


def _es_no_localizada(mensaje: dict) -> bool:
    return "NO LOCALIZADA" in str(mensaje.get("tipo", "")).upper()


def _esc(valor) -> str:
    """Escapa un valor para incrustarlo en el cuerpo HTML del correo."""
    return _escape_html(str(valor or ""), quote=False)


def _linea_contrato(numero: int, mensaje: dict) -> str:
    """Devuelve un contrato como item ``<li>`` de la lista HTML del cuerpo."""
    contrato = _esc(str(mensaje.get("contrato", "")).strip())
    linea = (
        f"<li><strong>{numero}. Contrato {contrato}</strong> — {_esc(mensaje.get('contratista', ''))} "
        f"— generado el día {_esc(mensaje.get('fecha_contrato', ''))}, "
        f"{_esc(mensaje.get('dias_transcurridos', ''))} "
        f"— Objeto: {_esc(mensaje.get('objeto', ''))}"
    )
    #faltantes = str(mensaje.get("listado_faltantes", "")).strip()
    #if not _es_no_localizada(mensaje) and faltantes:
       # linea += f"<br><strong>Documentos pendientes:</strong> {_esc(faltantes)}"
    return linea + "</li>"


def _lista_contratos_html(mensajes: list, inicio: int = 1) -> str:
    """Envuelve las líneas de contrato en una lista ordenada HTML."""
    items = "\n".join(
        _linea_contrato(numero, mensaje)
        for numero, mensaje in enumerate(mensajes, start=inicio)
    )
    if inicio > 1:
        return f"<ol start=\"{inicio}\">\n{items}\n</ol>"
    return f"<ol>\n{items}\n</ol>"


def _cuerpo_power_automate(mensajes: list) -> str:
    """Arma el cuerpo en HTML (un saludo, un marco y un solo cierre).

    Power Automate (conector Outlook, ``Is HTML = Sí``) ignora los saltos de
    línea del texto plano, por eso el cuerpo usa ``<p>`` y listas ``<ol>``.
    """
    primero = mensajes[0]
    ordenador = _esc(str(primero.get("para", "")).strip())
    marco = (
        f"En el marco del seguimiento contractual del mes de {_esc(primero.get('mes_documento', ''))} "
        f"de 2026 (2026-{_esc(primero.get('mes_numero', ''))}), le recordamos de manera atenta "
    )
    no_localizados = [m for m in mensajes if _es_no_localizada(m)]
    pendientes = [m for m in mensajes if not _es_no_localizada(m)]

    if len(mensajes) == 1:
        if no_localizados:
            intro = (
                "que el siguiente expediente contractual a su cargo aún no ha sido cargado "
                "en el repositorio institucional Alfresco:"
            )
        else:
            intro = (
                "que el expediente contractual ya fue localizado en Alfresco, pero aún presenta "
                "documentación pendiente:"
            )
        detalle = _lista_contratos_html(mensajes)
    elif no_localizados and not pendientes:
        intro = (
            "que los siguientes expedientes contractuales a su cargo aún no han sido cargados "
            "en el repositorio institucional Alfresco:"
        )
        detalle = _lista_contratos_html(no_localizados)
    elif pendientes and not no_localizados:
        intro = (
            "que los siguientes expedientes ya fueron localizados en Alfresco, pero aún presentan "
            "documentación pendiente:"
        )
        detalle = _lista_contratos_html(pendientes)
    else:
        intro = (
            "que los siguientes expedientes contractuales a su cargo presentan novedades "
            "en el repositorio institucional Alfresco:"
        )
        bloques = [
            "<p><strong>Expedientes no cargados en Alfresco:</strong></p>",
            _lista_contratos_html(no_localizados),
            "<p><strong>Expedientes localizados con documentación pendiente:</strong></p>",
            _lista_contratos_html(pendientes, inicio=len(no_localizados) + 1),
        ]
        detalle = "\n".join(bloques)

    return "\n".join([
        f"<p>Cordial saludo, {ordenador}.</p>",
        f"<p>{marco}{intro}</p>",
        detalle,
        f"<p>{_CIERRE_CORREO_HTML_1}</p>",
        f"<p>{_CIERRE_CORREO_HTML_2}</p>",
    ])


def _fecha_contrato(valor):
    """Convierte las fechas del consolidado a ``date`` cuando es posible."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor or "").strip()
    for formato in ("%Y/%m/%d", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto[:10], formato).date()
        except ValueError:
            continue
    return None


def _datos_fecha(fecha_contrato, fecha_documento: str = "") -> dict:
    """Obtiene el texto de fecha, el mes del documento y días transcurridos."""
    fecha = _fecha_contrato(fecha_contrato)
    fecha_revision = _fecha_contrato(fecha_documento) or date.today()
    mes_numero = f"{fecha_revision.month:02d}"
    return {
        "fecha_contrato": fecha.strftime("%Y/%m/%d") if fecha else str(fecha_contrato or ""),
        "mes_documento": (
            ("enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
             "septiembre", "octubre", "noviembre", "diciembre")[fecha_revision.month - 1]
        ),
        "mes_numero": mes_numero,
        "dias_transcurridos": (date.today() - fecha).days if fecha else "",
    }


def generar_borradores(
    contratos: list,
    ruta: str,
    ruta_diccionario: str = "",
    fecha_revision: str = "",
    fecha_limite: str = "",
    unidad_por_contrato: dict = None,
    ruta_excel: str = "",
) -> dict:
    """Carga plantillas, construye los borradores y los guarda.

    Devuelve ``{ruta, total, por_tipo}``.
    """
    plantillas = cargar_plantillas(ruta_diccionario)
    mensajes = construir_mensajes(
        contratos,
        plantillas,
        fecha_revision=fecha_revision,
        fecha_limite=fecha_limite,
        unidad_por_contrato=unidad_por_contrato,
    )
    escribir_borradores(mensajes, ruta)
    ruta_excel = ruta_excel or os.path.splitext(ruta)[0] + ".xlsx"
    ruta_excel_efectiva = escribir_borradores_excel(mensajes, ruta_excel)

    por_tipo = {}
    for mensaje in mensajes:
        por_tipo[mensaje["tipo"]] = por_tipo.get(mensaje["tipo"], 0) + 1
    return {
        "ruta": ruta,
        "ruta_excel": ruta_excel_efectiva,
        "total": len(mensajes),
        "por_tipo": por_tipo,
    }
