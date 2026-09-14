# -*- coding: utf-8 -*-
"""
Informe HTML del servicio «Auditoría Documental Automatizada».

Muestra el estado actual de los contratos procesados, clasificados en dos
módulos desplegables:

* **Con carpeta en Alfresco** — con el detalle de documentos **faltantes**
  (en un desplegable) y cuántos lleva del total de documentos obligatorios.
* **Sin carpeta en Alfresco** — en rojo, indicando que necesita carpeta.

Incluye un resumen con los procesados del día y el porcentaje de encontrados /
no encontrados, y una tabla general de contratos que se puede **descargar en
Excel** (el .xlsx se genera en Python y viaja embebido en el HTML, sin recursos
externos).
"""

import base64
import html
import io
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from utils.logger import configurar_logger
from .diccionario import ordenar_etapas

logger = configurar_logger("informe_html")

ENCONTRADO = "ENCONTRADO"
FALTANTE = "FALTANTE"
NO_APLICA = "NO APLICA"

_MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_ESTILOS = """
:root{--bg:#eef1f4;--panel:#fff;--ink:#1b2733;--muted:#64748b;--line:#d7dee6;
--accent:#145c56;--ok:#1e7d4f;--ok-soft:#e4f4ea;--warn:#b7791f;--warn-soft:#fdf3dc;
--bad:#b42318;--bad-soft:#fde8e6;--gray:#64748b;--gray-soft:#eef1f4;}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--ink);
font-family:Segoe UI,Roboto,Helvetica,Arial,sans-serif;font-size:14px;line-height:1.45}
header{background:var(--accent);color:#fff;padding:22px 26px}
header h1{margin:0 0 4px;font-size:21px}
header p{margin:0;opacity:.9;font-size:13px}
.wrap{max-width:1150px;margin:0 auto;padding:20px 26px 60px}
.cards{display:flex;gap:12px;flex-wrap:wrap;margin:18px 0}
.card{flex:1 1 140px;background:var(--panel);border:1px solid var(--line);border-radius:10px;padding:12px 14px}
.card .n{font-size:22px;font-weight:700}
.card .l{font-size:12px;color:var(--muted);text-transform:uppercase;letter-spacing:.04em}
.card.ok .n{color:var(--ok)} .card.bad .n{color:var(--bad)}
h2{font-size:16px;margin:24px 0 10px}
h2 .pill{font-size:12px;font-weight:600;background:var(--gray-soft);color:var(--muted);
border-radius:999px;padding:2px 10px;margin-left:8px}
.barra{display:flex;gap:10px;align-items:center;flex-wrap:wrap;margin-bottom:10px}
.btn{display:inline-block;background:var(--accent);color:#fff;text-decoration:none;
border:0;border-radius:8px;padding:8px 14px;font-size:13px;cursor:pointer}
.btn:hover{filter:brightness(1.1)}
table{border-collapse:collapse;width:100%;background:var(--panel);border:1px solid var(--line);
border-radius:10px;overflow:hidden;font-size:13px}
th,td{border-bottom:1px solid var(--line);padding:7px 9px;text-align:left;vertical-align:top}
th{background:#f4f7f9;font-size:12px;text-transform:uppercase;letter-spacing:.03em;color:var(--muted)}
tr:last-child td{border-bottom:0}
details.modulo{background:var(--panel);border:1px solid var(--line);border-radius:12px;margin:14px 0;overflow:hidden}
details.modulo>summary{cursor:pointer;padding:13px 16px;display:flex;gap:10px;align-items:center;flex-wrap:wrap;
font-size:16px;font-weight:700;background:#f4f7f9}
details.modulo>summary::-webkit-details-marker{display:none}
details.modulo .mod-body{padding:12px 14px}
details.contrato{background:var(--panel);border:1px solid var(--line);border-radius:10px;margin-bottom:10px;overflow:hidden}
details.contrato>summary{cursor:pointer;padding:11px 14px;display:flex;gap:10px;align-items:center;flex-wrap:wrap}
details.contrato>summary::-webkit-details-marker{display:none}
details.sin>summary{background:var(--bad-soft)}
details.sin{border-color:#f3c9c4}
summary .cid{font-weight:700}
summary .meta{color:var(--muted);font-size:12px}
.pill2{background:var(--gray-soft);color:var(--muted);border-radius:999px;padding:2px 9px;font-size:12px}
.badge{margin-left:auto;padding:3px 10px;border-radius:999px;font-size:12px;font-weight:600}
.b-ok{background:var(--ok-soft);color:var(--ok)}
.b-warn{background:var(--warn-soft);color:var(--warn)}
.b-bad{background:var(--bad-soft);color:var(--bad)}
.b-gray{background:var(--gray-soft);color:var(--gray)}
.necesita{color:var(--bad);font-weight:700}
.cuerpo{padding:0 14px 14px;border-top:1px solid var(--line)}
.fin{display:flex;gap:16px;flex-wrap:wrap;color:var(--muted);font-size:12px;margin:10px 0}
details.sub{border:1px solid var(--line);border-radius:8px;margin-top:8px;overflow:hidden}
details.sub>summary{cursor:pointer;padding:8px 12px;font-weight:600;font-size:13px}
details.sub .lista{padding:0 12px 10px}
.etapa{font-size:12px;font-weight:700;color:var(--muted);margin:8px 0 4px}
ul.docs{list-style:none;margin:0;padding:0}
ul.docs li{border:1px solid var(--line);border-radius:8px;padding:6px 9px;margin-bottom:5px;font-size:13px}
ul.docs li .fmt{font-weight:600}
ul.docs li .arch{display:block;color:var(--muted);font-size:12px;margin-top:2px;word-break:break-all}
li.f{border-color:#f3c9c4;background:#fff7f6}
li.n{background:#f7f9fb;color:var(--muted)}
.vacio{color:var(--muted);font-size:13px}
a{color:var(--accent)}
footer{color:var(--muted);font-size:12px;text-align:center;padding:20px}
@media print{
  header{-webkit-print-color-adjust:exact;print-color-adjust:exact}
  .barra,.btn{display:none}
  .wrap{max-width:none;padding:0}
  details{break-inside:avoid}
  details>summary{background:#f4f7f9}
  .card{break-inside:avoid}
}
"""

_CABECERAS_EXCEL = [
    "CONTRATO", "CLASE", "CONTRATISTA", "VALOR", "FECHA INICIO", "FECHA FIN",
    "ORDENADOR", "ESTADO", "DOCS HALLADOS", "DOCS REQUERIDOS",
    "DOCS FALTANTES", "% COMPLETITUD", "CARPETA ALFRESCO", "CARPETA",
]


def _e(valor) -> str:
    return html.escape(str(valor if valor is not None else ""))


def _num(valor, defecto=0) -> int:
    try:
        return int(float(str(valor).strip() or defecto))
    except (TypeError, ValueError):
        return defecto


def _estado(resultado: dict) -> str:
    if str(resultado.get("estado_alfresco", "")).upper() != "ENCONTRADA":
        return "No encontrada"
    return "Completo" if _num(resultado.get("faltantes")) == 0 else "Incompleto"


def _badge(estado: str, pct) -> str:
    if estado == "Completo":
        return '<span class="badge b-ok">Completo</span>'
    if estado == "Incompleto":
        clase = "b-ok" if (pct or 0) >= 80 else ("b-warn" if (pct or 0) >= 50 else "b-bad")
        return f'<span class="badge {clase}">Incompleto · {_e(pct)}%</span>'
    if estado == "No encontrada":
        return '<span class="badge b-bad">No encontrada</span>'
    return f'<span class="badge b-gray">{_e(estado)}</span>'


def _por_estado(filas: list, estado: str) -> list:
    return [f for f in (filas or []) if str(f.get("ESTADO", "")).upper() == estado]


def _agrupar_por_etapa(filas: list) -> dict:
    grupos = {}
    for fila in filas:
        etapa = fila.get("ETAPA") or "Sin etapa"
        grupos.setdefault(etapa, []).append(fila)
    return grupos


def _lista_documentos(filas: list, con_archivo: bool, clase_li: str) -> str:
    if not filas:
        return '<p class="vacio">—</p>'
    grupos = _agrupar_por_etapa(filas)
    partes = []
    for etapa in ordenar_etapas(grupos.keys()):
        partes.append(f'<div class="etapa">{_e(etapa)}</div><ul class="docs">')
        for fila in grupos[etapa]:
            formato = _e(fila.get("FORMATO", "")).strip()
            documento = _e(fila.get("DOCUMENTO", "")).strip()
            archivo = ""
            if con_archivo and str(fila.get("ARCHIVO", "")).strip():
                archivo = f'<span class="arch">{_e(fila.get("ARCHIVO", ""))}</span>'
            partes.append(
                f'<li class="{clase_li}"><span class="fmt">{formato}</span> '
                f'<span class="doc">{documento}</span>{archivo}</li>'
            )
        partes.append("</ul>")
    return "".join(partes)


def _sub(nombre: str, filas: list, con_archivo: bool, clase_li: str, abierto: bool = False) -> str:
    if not filas:
        return ""
    open_attr = " open" if abierto else ""
    return (
        f'<details class="sub"{open_attr}><summary>{_e(nombre)} ({len(filas)})</summary>'
        f'<div class="lista">{_lista_documentos(filas, con_archivo, clase_li)}</div></details>'
    )


def _info_financiera(resultado: dict) -> str:
    campos = [
        ("Contratista", resultado.get("contratista", "")),
        ("Valor", resultado.get("valor", "")),
        ("Inicio", resultado.get("fecha_inicio", "")),
        ("Fin", resultado.get("fecha_fin", "")),
        ("Ordenador", resultado.get("supervisor", "")),
        ("Unidad", resultado.get("unidad", "")),
    ]
    partes = [f"<span>{_e(k)}: {_e(v)}</span>" for k, v in campos if str(v).strip()]
    return f'<div class="fin">{"".join(partes)}</div>' if partes else ""


def _tarjeta_con_carpeta(resultado: dict) -> str:
    filas = resultado.get("filas") or []
    encontrados = _por_estado(filas, ENCONTRADO)
    faltantes = _por_estado(filas, FALTANTE)
    no_aplica = _por_estado(filas, NO_APLICA)
    esperados = _num(resultado.get("esperados"))
    hallados = _num(resultado.get("encontrados"))
    archivos = _num(resultado.get("cantidad_archivos"))
    pct = resultado.get("pct_cumpl", 0)
    node_id = resultado.get("node_id", "")
    enlace = (
        f'<a href="https://gesdoc.uis.edu.co/share/page/folder-details?nodeRef='
        f'workspace://SpacesStore/{_e(node_id)}" target="_blank">Abrir carpeta en Alfresco</a>'
        if node_id else ""
    )
    return f"""
<details class="contrato">
  <summary>
    <span class="cid">{_e(resultado.get('contrato',''))}</span>
    <span class="meta">Clase {_e(resultado.get('cod',''))} · {_e(resultado.get('carpeta',''))}</span>
    <span class="pill2">Obligatorios {hallados}/{esperados}</span>
    <span class="pill2">Archivos en carpeta {archivos}</span>
    {_badge(_estado(resultado), pct)}
  </summary>
  <div class="cuerpo">
    {_info_financiera(resultado)}
    {_sub("Documentos faltantes", faltantes, False, "f", abierto=True)}
    {_sub("Documentos encontrados", encontrados, True, "", abierto=False)}
    {_sub("No aplica", no_aplica, False, "n", abierto=False)}
    <p class="meta">{enlace}</p>
  </div>
</details>"""


def _tarjeta_sin_carpeta(resultado: dict) -> str:
    return f"""
<details class="contrato sin">
  <summary>
    <span class="cid">{_e(resultado.get('contrato',''))}</span>
    <span class="meta">Clase {_e(resultado.get('cod',''))}</span>
    <span class="necesita">Necesita carpeta en Alfresco</span>
    <span class="meta">{_e(resultado.get('contratista',''))}</span>
  </summary>
  <div class="cuerpo">
    {_info_financiera(resultado)}
    <p class="vacio">No se encontró la carpeta del expediente en Alfresco; se debe crear o ubicar para poder verificar su documentación.</p>
  </div>
</details>"""


def _tabla_resumen(resultados: list) -> str:
    cabeceras = [
        "CONTRATO", "CLASE", "CONTRATISTA", "VALOR", "INICIO", "FIN",
        "ORDENADOR", "ESTADO", "DOCS", "%", "ALFRESCO",
    ]
    filas_html = []
    for r in resultados:
        estado = _estado(r)
        alfre = "Sí" if str(r.get("estado_alfresco", "")).upper() == "ENCONTRADA" else "No"
        filas_html.append(
            "<tr>"
            f"<td>{_e(r.get('contrato',''))}</td>"
            f"<td>{_e(r.get('cod',''))}</td>"
            f"<td>{_e(r.get('contratista',''))}</td>"
            f"<td>{_e(r.get('valor',''))}</td>"
            f"<td>{_e(r.get('fecha_inicio',''))}</td>"
            f"<td>{_e(r.get('fecha_fin',''))}</td>"
            f"<td>{_e(r.get('supervisor',''))}</td>"
            f"<td>{_e(estado)}</td>"
            f"<td>{_num(r.get('encontrados'))}/{_num(r.get('esperados'))}</td>"
            f"<td>{_e(r.get('pct_cumpl',0))}%</td>"
            f"<td>{alfre}</td>"
            "</tr>"
        )
    encabezados = "".join(f"<th>{_e(c)}</th>" for c in cabeceras)
    return (
        '<table id="tabla-resumen"><thead><tr>' + encabezados + "</tr></thead><tbody>"
        + "".join(filas_html) + "</tbody></table>"
    )


def _excel_resumen_b64(resultados: list) -> str:
    """Genera el Excel de la tabla resumen y lo devuelve en base64."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Contratos"

    hoja.append(_CABECERAS_EXCEL)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F4E78")

    for r in resultados:
        encontrada = str(r.get("estado_alfresco", "")).upper() == "ENCONTRADA"
        hoja.append([
            r.get("contrato", ""),
            r.get("cod", ""),
            r.get("contratista", ""),
            r.get("valor", ""),
            r.get("fecha_inicio", ""),
            r.get("fecha_fin", ""),
            r.get("supervisor", ""),
            _estado(r),
            _num(r.get("encontrados")),
            _num(r.get("esperados")),
            _num(r.get("faltantes")),
            f"{r.get('pct_cumpl', 0)}%",
            "Sí" if encontrada else "No",
            r.get("carpeta", ""),
        ])

    for columna, ancho in enumerate([16, 8, 30, 15, 12, 12, 26, 14, 12, 14, 12, 12, 10, 24], start=1):
        hoja.column_dimensions[get_column_letter(columna)].width = ancho
    hoja.freeze_panes = "A2"

    buffer = io.BytesIO()
    libro.save(buffer)
    return base64.b64encode(buffer.getvalue()).decode("ascii")


def generar_informe(resultados: list, ruta: str, periodo: str = "", fecha_revision: str = "") -> str:
    """Genera el informe HTML y devuelve la ruta."""
    resultados = resultados or []
    total = len(resultados)
    con_carpeta = [r for r in resultados if str(r.get("estado_alfresco", "")).upper() == "ENCONTRADA"]
    sin_carpeta = [r for r in resultados if str(r.get("estado_alfresco", "")).upper() != "ENCONTRADA"]
    total_encontrados = sum(_num(r.get("encontrados")) for r in resultados)
    total_faltantes = sum(_num(r.get("faltantes")) for r in resultados)
    pct_encontrados = round(len(con_carpeta) * 100 / total, 1) if total else 0
    pct_sin = round(len(sin_carpeta) * 100 / total, 1) if total else 0
    pct_promedio = round(sum(float(r.get("pct_cumpl", 0) or 0) for r in resultados) / total, 1) if total else 0

    tarjetas_con = "".join(_tarjeta_con_carpeta(r) for r in con_carpeta)
    tarjetas_sin = "".join(_tarjeta_sin_carpeta(r) for r in sin_carpeta)
    if not con_carpeta:
        tarjetas_con = '<p class="vacio">Ningún contrato con carpeta.</p>'
    if not sin_carpeta:
        tarjetas_sin = '<p class="vacio">Todos los contratos tienen carpeta.</p>'

    excel_b64 = _excel_resumen_b64(resultados)
    nombre_excel = f"contratos_{periodo}.xlsx" if periodo else "contratos_auditoria.xlsx"

    documento = f"""<!doctype html>
<html lang="es"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Auditoría Documental Automatizada</title>
<style>{_ESTILOS}</style></head>
<body>
<header>
  <h1>Auditoría Documental Automatizada</h1>
  <p>Estado actual de los contratos procesados y su documentación.
  {'Periodo: ' + _e(periodo) + ' · ' if periodo else ''}Revisión: {_e(fecha_revision or '—')}</p>
</header>
<div class="wrap">
  <div class="cards">
    <div class="card"><div class="n">{total}</div><div class="l">Contratos procesados</div></div>
    <div class="card ok"><div class="n">{len(con_carpeta)}</div><div class="l">Con carpeta ({pct_encontrados}%)</div></div>
    <div class="card bad"><div class="n">{len(sin_carpeta)}</div><div class="l">Sin carpeta ({pct_sin}%)</div></div>
    <div class="card"><div class="n">{total_encontrados}</div><div class="l">Documentos hallados</div></div>
    <div class="card"><div class="n">{total_faltantes}</div><div class="l">Documentos faltantes</div></div>
    <div class="card"><div class="n">{pct_promedio}%</div><div class="l">Completitud promedio</div></div>
  </div>

  <h2>Resumen de contratos <span class="pill">visualiza y descarga</span></h2>
  <div class="barra">
    <a class="btn" download="{_e(nombre_excel)}" href="data:{_MIME_XLSX};base64,{excel_b64}">Descargar Excel</a>
    <a class="btn" href="javascript:imprimir()">Descargar PDF</a>
    <span class="vacio">Descarga la tabla en Excel o guarda el informe completo en PDF.</span>
  </div>
  {_tabla_resumen(resultados)}

  <details class="modulo" open>
    <summary>Contratos con carpeta en Alfresco <span class="pill">{len(con_carpeta)}</span></summary>
    <div class="mod-body">
      <p class="vacio">Despliega cada contrato para ver los documentos que le faltan y los que ya tiene.</p>
      {tarjetas_con}
    </div>
  </details>

  <details class="modulo" open>
    <summary>Contratos sin carpeta en Alfresco <span class="pill">{len(sin_carpeta)}</span></summary>
    <div class="mod-body">
      <p class="vacio">Estos contratos requieren que se cree o ubique su carpeta en el repositorio.</p>
      {tarjetas_sin}
    </div>
  </details>
</div>
<footer>Generado por el servicio «Auditoría Documental Automatizada».</footer>
<script>
function imprimir(){{
  document.querySelectorAll('details').forEach(function(d){{ d.open = true; }});
  window.print();
}}
</script>
</body></html>"""

    os.makedirs(os.path.dirname(ruta) or ".", exist_ok=True)
    with open(ruta, "w", encoding="utf-8") as archivo:
        archivo.write(documento)
    logger.info("Informe HTML guardado: %s (%d contratos)", ruta, total)
    return ruta
